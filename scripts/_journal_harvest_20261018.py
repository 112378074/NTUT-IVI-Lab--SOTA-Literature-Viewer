#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Journal harvest update — 2026-10-18
Weekly journal-targeted harvest; adds new AD candidates found this week.

Search scope: IEEE TPAMI, IEEE TII, IEEE TNNLS, Pattern Recognition, IJCV,
  IEEE TIP, IEEE TCSVT, Computers in Industry, Information Fusion,
  Knowledge-Based Systems, Information Sciences, Neural Networks.

Findings this run:
  — No NEW confirmed target-journal papers found (IJCV, TII, TPAMI, etc.)
    this week. ResAD++ (arXiv 2509.23741, IJCV 2026) was already staged
    as ⏳ Pending in a prior harvest; no metric update possible (arXiv blocked,
    GitHub only has ResAD predecessor).
  — 2 new October 2026 arXiv papers found (not in workbook):
    * CORE (arXiv 2610.19610, Oct 14, 2026) — anomaly synthesis + LMM-based
      prompt optimization; MVTec AD + VisA. arXiv HTML blocked, no GitHub.
      Staged ⏳ Pending.
    * InCLAD (arXiv 2610.15032, Oct 10, 2026) — continual learning benchmark
      for IAD; MVTec AD (15 tasks) + VisA (12 tasks). NON-STANDARD PROTOCOL.
      arXiv HTML blocked. Staged ⏳ Pending.
  — Papers confirmed already in workbook (skipped): ResAD++ (2509.23741),
    Anchor and Adapt (2610.07016), VAND 4.0 (2610.04392).
  — Not applicable: Cross-Scenario Defect Detection (2610.19540, IEEE ICME 2026
    — cross-domain detection, not AD on MVTec/VisA family for this workbook);
    Certification limits (2610.00236 — analysis/theory paper, no method rows);
    MRAD (arXiv 2602.00522, ICLR 2026 — conference, not target journal).
  — arXiv.org still blocked by proxy; no GitHub repos found for new candidates.
  — All new entries staged as ⏳ Pending with blank metrics.
"""

from pathlib import Path
import sys

try:
    import openpyxl
except ImportError:
    sys.path.insert(0, '/tmp/pkgs')
    import openpyxl

AD_XLSX = Path(__file__).resolve().parent.parent / 'AnomalyDetection_Papers_Summary_v10_20260425.xlsx'

# ---------------------------------------------------------------------------
# Row data — column order (15 cols):
# (資料集, 類別, 方法, 作者, 發表, 年月, 狀態,
#  I-AUROC, P-AUROC, P-AP, P-PRO, FPS, 備註, 連結, GitHub)
# ---------------------------------------------------------------------------

# --- 1. CORE — arXiv 2610.19610, Oct 14 2026, anomaly synthesis + LMM ---

CORE_OVERVIEW = (
    'MVTec AD, VisA (異常合成框架; 評估下游 AD 效果)',
    '基於異常合成 (Anomaly Synthesis-based) / 基礎模型/大型多模態模型 (Foundation Model / LMM-based) '
    '/ 自監督提示優化 (Self-Supervised Prompt Optimization)',
    'CORE (Contrast-and-Reflect for Self-Supervised Prompt Optimization in Anomaly Synthesis)',
    'Haoxiang Rao, Fang Zhao, Zhao Wang, Weijia Li, Chenyang Si, Caifeng Shan',
    'arXiv 2610.19610 (2026-10-14)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2610.19610 (Oct 14, 2026); arXiv HTML blocked, no GitHub repo found. '
    'Training-free self-supervised prompt optimization framework for anomaly synthesis using '
    'large multimodal models (LMMs). Proposes an Anomaly Context Grounding (ACG) module to '
    'generate location-conditioned prompts, and a Contrast-and-Reflect (CoR) step that '
    'iteratively refines prompts by contrasting synthesized anomalies against normal regions '
    'in a self-supervised loop. Covers both zero-shot and one-shot anomaly synthesis settings. '
    'Evaluates by measuring downstream AD performance gains on MVTec-AD and VisA benchmarks. '
    'Method improves synthesis quality for downstream detectors; specific I-AUROC/P-AUROC '
    'values for the downstream detector not accessible (arXiv blocked). '
    'Journal harvest 2026-10-18. [基於異常合成 / 基礎模型/LMM / 自監督提示優化]',
    'https://arxiv.org/abs/2610.19610',
    None,
)

CORE_MVTEC = (
    'MVTec AD (異常合成框架; 評估下游 AD 效果)',
    '基於異常合成 (Anomaly Synthesis-based) / 基礎模型/LMM (Foundation Model / LMM-based) '
    '/ 自監督提示優化 (Self-Supervised Prompt Optimization)',
    'CORE (Contrast-and-Reflect for Self-Supervised Prompt Optimization in Anomaly Synthesis)',
    'Haoxiang Rao, Fang Zhao, Zhao Wang, Weijia Li, Chenyang Si, Caifeng Shan',
    'arXiv 2610.19610 (2026-10-14)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — anomaly synthesis framework; downstream AD metrics not accessible (arXiv blocked). '
    '[基於異常合成 / 基礎模型/LMM / 自監督提示優化]',
    'https://arxiv.org/abs/2610.19610',
    None,
)

CORE_VISA = (
    'VisA (異常合成框架; 評估下游 AD 效果)',
    '基於異常合成 (Anomaly Synthesis-based) / 基礎模型/LMM (Foundation Model / LMM-based) '
    '/ 自監督提示優化 (Self-Supervised Prompt Optimization)',
    'CORE (Contrast-and-Reflect for Self-Supervised Prompt Optimization in Anomaly Synthesis)',
    'Haoxiang Rao, Fang Zhao, Zhao Wang, Weijia Li, Chenyang Si, Caifeng Shan',
    'arXiv 2610.19610 (2026-10-14)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — anomaly synthesis framework; downstream AD metrics not accessible (arXiv blocked). '
    '[基於異常合成 / 基礎模型/LMM / 自監督提示優化]',
    'https://arxiv.org/abs/2610.19610',
    None,
)

# --- 2. InCLAD — arXiv 2610.15032, Oct 10 2026, continual learning benchmark ---

INCLAD_OVERVIEW = (
    'MVTec AD (15 tasks), VisA (12 tasks) (持續學習基準; NON-STANDARD PROTOCOL — 持續/漸進學習)',
    '基準/評測論文 (Benchmark Paper) / 持續學習 (Continual Learning) / 工業視覺異常偵測',
    'InCLAD (A Continual Learning Benchmark for Industrial Visual Anomaly Detection)',
    'Łukasz Marcjan, Dorota Wilk-Kołodziejczyk, Roberto Corizzo',
    'arXiv 2610.15032 (2026-10-10)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2610.15032 (Oct 10, 2026); arXiv HTML blocked. '
    'Benchmark paper defining continual learning evaluation protocols for industrial VAD. '
    'Task-incremental setup: each industrial component is a separate task; training is '
    'semi-supervised (normal images only). Defines three continual scenarios: Random, '
    'Easy-to-Hard, and Hard-to-Easy task orderings, plus cross-dataset scenarios. '
    'Key challenge: catastrophic forgetting. Evaluates existing AD methods (standard and '
    'continual) on MVTec AD (15 tasks) and VisA (12 tasks). '
    'NOT a method paper — catalog reference entry for benchmark evaluation protocol. '
    'NON-STANDARD PROTOCOL (continual/incremental learning). '
    'Journal harvest 2026-10-18. [基準/評測 / 持續學習]',
    'https://arxiv.org/abs/2610.15032',
    None,
)

INCLAD_MVTEC = (
    'MVTec AD (15 tasks, 持續學習基準; NON-STANDARD PROTOCOL — 持續/漸進學習)',
    '基準/評測論文 (Benchmark Paper) / 持續學習 (Continual Learning)',
    'InCLAD (A Continual Learning Benchmark for Industrial Visual Anomaly Detection)',
    'Łukasz Marcjan, Dorota Wilk-Kołodziejczyk, Roberto Corizzo',
    'arXiv 2610.15032 (2026-10-10)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — continual learning benchmark on MVTec AD 15 tasks. arXiv HTML blocked; '
    'per-method forgetting/accuracy metrics not accessible. NOT comparable to standard '
    'unsupervised AD. [基準/評測 / 持續學習]',
    'https://arxiv.org/abs/2610.15032',
    None,
)

INCLAD_VISA = (
    'VisA (12 tasks, 持續學習基準; NON-STANDARD PROTOCOL — 持續/漸進學習)',
    '基準/評測論文 (Benchmark Paper) / 持續學習 (Continual Learning)',
    'InCLAD (A Continual Learning Benchmark for Industrial Visual Anomaly Detection)',
    'Łukasz Marcjan, Dorota Wilk-Kołodziejczyk, Roberto Corizzo',
    'arXiv 2610.15032 (2026-10-10)',
    '2026-10',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — continual learning benchmark on VisA 12 tasks. arXiv HTML blocked; '
    'metrics not accessible. NOT comparable to standard unsupervised AD. [基準/評測 / 持續學習]',
    'https://arxiv.org/abs/2610.15032',
    None,
)


# ---------------------------------------------------------------------------
# Insert rows
# ---------------------------------------------------------------------------

def sheet_has_link(ws, link):
    link_lower = link.lower()
    for row in ws.iter_rows(values_only=True):
        if len(row) > 13 and row[13] and str(row[13]).lower() == link_lower:
            return True
    return False


def safe_append(ws, row_tuple):
    link = row_tuple[13]
    if link and sheet_has_link(ws, link):
        print(f'  SKIP (dupe): {row_tuple[2][:55]} → {ws.title}')
        return False
    ws.append(row_tuple)
    print(f'  ADD: {row_tuple[2][:55]} → {ws.title}')
    return True


def main():
    wb = openpyxl.load_workbook(str(AD_XLSX))
    ws_all   = wb['総覧 All Papers'] if '総覧 All Papers' in wb.sheetnames else wb['總覽 All Papers']
    ws_mvtec = wb['MVTec AD']
    ws_visa  = wb['VisA']

    print('=== Adding CORE ===')
    safe_append(ws_all,   CORE_OVERVIEW)
    safe_append(ws_mvtec, CORE_MVTEC)
    safe_append(ws_visa,  CORE_VISA)

    print('=== Adding InCLAD ===')
    safe_append(ws_all,   INCLAD_OVERVIEW)
    safe_append(ws_mvtec, INCLAD_MVTEC)
    safe_append(ws_visa,  INCLAD_VISA)

    wb.save(str(AD_XLSX))
    print(f'\n✅ Saved: {AD_XLSX}')
    print('New entries: 2 papers, 6 sheet rows total (overview + per-dataset)')


if __name__ == '__main__':
    main()
