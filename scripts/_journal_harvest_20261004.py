#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Journal harvest update — 2026-10-04
Weekly journal-targeted harvest; adds new AD candidates found this week.

Search scope: IEEE TPAMI, IEEE TII, IEEE TNNLS, Pattern Recognition, IJCV,
  IEEE TIP, IEEE TCSVT, Computers in Industry, Information Fusion,
  Knowledge-Based Systems, Information Sciences, Neural Networks.

Findings this run:
  — 1 new CONFIRMED journal paper found:
    * GPAD (arXiv 2603.22757) accepted in IEEE TCSVT — 3D/RGB-D, MVTec 3D-AD + Eyecandies
  — 3 new arXiv preprints staged as ⏳ Pending (missed in prior harvests):
    * GaugeDefect (2609.13282, Sep 2026) — standard unsupervised AD, MVTec AD/VisA/Real-IAD
      web-snippet AUPRO: MVTec 96.5 / VisA 95.1 / Real-IAD 94.4 (NOT verified — arXiv blocked)
    * Training-Free Logical+Structural AD via Calibrated Fusion (2609.05091, Sep 2026)
      — training-free protocol, MVTec LOCO; web-snippet avg AUROC 92.5 (NOT verified)
    * LUMIN (2609.04775, Sep 2026) — zero-shot cross-domain protocol,
      MVTec-AD/VisA/BTAD/KSDD/Real-IAD (no metrics found)
  — arXiv.org blocked by proxy; no GitHub repos found for any new candidates.
    All new entries staged as ⏳ Pending with blank metrics.
  — No new October 2026 (2610.xxxxx) papers indexed yet as of 2026-10-04.
"""

from pathlib import Path
import openpyxl

AD_XLSX = Path(__file__).resolve().parent.parent / 'AnomalyDetection_Papers_Summary_v10_20260425.xlsx'

# ---------------------------------------------------------------------------
# Row data — column order (15 cols):
# (資料集, 類別, 方法, 作者, 發表, 年月, 狀態,
#  I-AUROC, P-AUROC, P-AP, P-PRO, FPS, 備註, 連結, GitHub)
# ---------------------------------------------------------------------------

# --- 1. GPAD — IEEE TCSVT 2026, MVTec 3D-AD + Eyecandies ---

GPAD_OVERVIEW = (
    'MVTec 3D-AD, Eyecandies',
    '基於 3D/RGB-D (3D/RGB-D-based) / 幾何先驗 (Geometric Prior) / 點雲 (Point Cloud)',
    'GPAD (Multimodal Industrial Anomaly Detection via Geometric Prior)',
    'Min Li, Jinghui He, Gang Li, Jiachen Li, Jin Wan, Delong Han',
    'IEEE TCSVT 2026 (accepted) / arXiv 2603.22757 (2026-03-24)',
    '2026-03',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2603.22757 (Mar 24, 2026); accepted IEEE Transactions on Circuits and '
    'Systems for Video Technology (TCSVT). arXiv HTML blocked, no GitHub repo found. '
    'Proposes point cloud expert model with differential normal vector computation for fine-grained '
    'geometric feature extraction + two-stage fusion strategy (attention fusion + anomaly region '
    'segmentation based on geometric prior). PyTorch 1.13.1, RTX 4090, batch 8, 200 epochs. '
    '"Outperforms SOTA on MVTec 3D-AD and Eyecandies" (no specific AUROC accessible). '
    'Journal harvest 2026-10-04. [基於 3D/RGB-D / 幾何先驗 / 點雲]',
    'https://arxiv.org/abs/2603.22757',
    None,
)

GPAD_3D = (
    'MVTec 3D-AD',
    '基於 3D/RGB-D (3D/RGB-D-based) / 幾何先驗 (Geometric Prior) / 點雲 (Point Cloud)',
    'GPAD (Multimodal Industrial Anomaly Detection via Geometric Prior)',
    'Min Li, Jinghui He, Gang Li, Jiachen Li, Jin Wan, Delong Han',
    'IEEE TCSVT 2026 (accepted) / arXiv 2603.22757 (2026-03-24)',
    '2026-03',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked, no GitHub; "outperforms SOTA on MVTec 3D-AD". '
    'IEEE TCSVT accepted. [基於 3D/RGB-D / 幾何先驗]',
    'https://arxiv.org/abs/2603.22757',
    None,
)

# --- 2. GaugeDefect — arXiv 2609.13282, standard unsupervised AD ---

GAUGE_OVERVIEW = (
    'MVTec AD, VisA, Real-IAD',
    '基於表徵 (Representation-based) / 特徵傳輸曲率 (Feature Transport Curvature / Holonomy)',
    'GaugeDefect (Detecting Surface Anomalies by Curvature of Feature Transport)',
    'see arXiv 2609.13282',
    'arXiv 2609.13282 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.13282 (Sep 2026); arXiv HTML blocked, no GitHub repo found. '
    'Gauge-theoretic method: extracts spatial feature lattice from pretrained backbone, estimates '
    'local feature frames, measures feature-transport curvature (holonomy around closed loops); '
    'anomaly map from loop curvature calibrated with normal statistics. Gains clearest on '
    'localization-sensitive metrics (P-AP, AUPRO). '
    'Web-search snippet AUPRO (NOT verified from authoritative table): '
    'MVTec AD 96.5 / VisA 95.1 / Real-IAD 94.4. '
    'Journal harvest 2026-10-04. [基於表徵 / 特徵傳輸曲率]',
    'https://arxiv.org/abs/2609.13282',
    None,
)

GAUGE_MVTEC = (
    'MVTec AD',
    '基於表徵 (Representation-based) / 特徵傳輸曲率 (Feature Transport Curvature / Holonomy)',
    'GaugeDefect (Detecting Surface Anomalies by Curvature of Feature Transport)',
    'see arXiv 2609.13282',
    'arXiv 2609.13282 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-snippet AUPRO≈96.5 on MVTec AD (NOT verified). [基於表徵]',
    'https://arxiv.org/abs/2609.13282',
    None,
)

GAUGE_VISA = (
    'VisA',
    '基於表徵 (Representation-based) / 特徵傳輸曲率 (Feature Transport Curvature / Holonomy)',
    'GaugeDefect (Detecting Surface Anomalies by Curvature of Feature Transport)',
    'see arXiv 2609.13282',
    'arXiv 2609.13282 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-snippet AUPRO≈95.1 on VisA (NOT verified). [基於表徵]',
    'https://arxiv.org/abs/2609.13282',
    None,
)

# --- 3. Training-Free Logical+Structural AD via Calibrated Fusion (2609.05091) ---

LOGFREE_OVERVIEW = (
    'MVTec LOCO (training-free zero-shot protocol; NON-STANDARD PROTOCOL)',
    '基於基礎模型 (Foundation Model-based) / 訓練免費 (Training-Free) / 邏輯+結構異常 (Logical+Structural AD)',
    'CalibFusion-AD (Training-Free Logical and Structural Anomaly Detection via Calibrated Fusion)',
    'Changyi Li, Miao Yu, Kai Dong, Yu Xiao (Aalto University / Harbin Engineering University)',
    'arXiv 2609.05091 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.05091 (Sep 2026); arXiv HTML blocked, no GitHub repo found. '
    'Training-free approach that introduces counting ability for logical anomaly detection; '
    'normal-set calibration aligns heterogeneous anomaly cues (structural + logical) without '
    'additional training or part-level supervision. Best among training-free detectors on MVTec LOCO. '
    'Web-search snippet MVTec LOCO AUROC (NOT verified): logical 89.0 / structural 95.9 / avg 92.5. '
    'NON-STANDARD (training-free/zero-shot protocol; not comparable to trained standard AD). '
    'Journal harvest 2026-10-04. [基於基礎模型 / 訓練免費 / 邏輯+結構]',
    'https://arxiv.org/abs/2609.05091',
    None,
)

LOGFREE_LOCO = (
    'MVTec LOCO (training-free zero-shot protocol; NON-STANDARD PROTOCOL)',
    '基於基礎模型 (Foundation Model-based) / 訓練免費 (Training-Free) / 邏輯+結構異常 (Logical+Structural AD)',
    'CalibFusion-AD (Training-Free Logical and Structural Anomaly Detection via Calibrated Fusion)',
    'Changyi Li, Miao Yu, Kai Dong, Yu Xiao (Aalto / Harbin EUT)',
    'arXiv 2609.05091 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-snippet avg AUROC≈92.5 (logical 89.0, structural 95.9) '
    'on MVTec LOCO (NOT verified; training-free NON-STANDARD protocol). [基於基礎模型 / 訓練免費]',
    'https://arxiv.org/abs/2609.05091',
    None,
)

# --- 4. LUMIN — arXiv 2609.04775, zero-shot cross-domain ---

LUMIN_OVERVIEW = (
    'MVTec AD, VisA, BTAD, KSDD, Real-IAD (zero-shot cross-domain protocol; NON-STANDARD PROTOCOL)',
    '基於基礎模型 (Foundation Model-based) / CLIP + DINOv3 / 記憶庫 (Memory Bank) / 零樣本跨域 (Zero-Shot Cross-Domain)',
    'LUMIN (Lightweight Universal Manufacturing Inspection Network for Anomaly Detection)',
    'see arXiv 2609.04775',
    'arXiv 2609.04775 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.04775 (Sep 2026); arXiv HTML blocked, no GitHub repo found. '
    'Engineering efficiency focus: PSP (Plugin Sampler Pipeline) — adaptive memory bank sampling '
    'with zero backbone forward passes (341× faster than FPS); parallel memory bank similarity '
    'computation (reduces inference memory/latency by >95%); stratified pixel sampling (20× '
    'evaluation speedup). Zero-shot cross-domain protocol: trained on VisA, tested on MVTec-AD, '
    'BTAD, KSDD, Real-IAD without fine-tuning; VisA evaluated with model trained on MVTec-AD. '
    'Frozen CLIP + DINOv3 backbones, input 518×518 + 512×512, layers 6/12/18/24. '
    'NON-STANDARD (cross-domain zero-shot; not comparable to standard AD benchmarks). '
    'No dataset-specific metrics available. '
    'Journal harvest 2026-10-04. [基於基礎模型 / CLIP + DINOv3 / 記憶庫 / 零樣本跨域]',
    'https://arxiv.org/abs/2609.04775',
    None,
)

LUMIN_MVTEC = (
    'MVTec AD (zero-shot cross-domain protocol; NON-STANDARD PROTOCOL)',
    '基於基礎模型 (Foundation Model-based) / CLIP + DINOv3 / 零樣本跨域 (Zero-Shot Cross-Domain)',
    'LUMIN (Lightweight Universal Manufacturing Inspection Network for Anomaly Detection)',
    'see arXiv 2609.04775',
    'arXiv 2609.04775 (2026-09)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; zero-shot cross-domain protocol (trained on VisA→tested '
    'on MVTec AD). No specific metrics found. NON-STANDARD. [基於基礎模型 / CLIP+DINOv3 / 零樣本跨域]',
    'https://arxiv.org/abs/2609.04775',
    None,
)

# ---------------------------------------------------------------------------
# Insert rows
# ---------------------------------------------------------------------------

def main():
    wb = openpyxl.load_workbook(str(AD_XLSX))
    ws_all  = wb['總覽 All Papers']
    ws_mvtec = wb['MVTec AD']
    ws_visa  = wb['VisA']
    ws_loco  = wb['MVTec LOCO']
    ws_3d    = wb['MVTec 3D']

    # ---- overview: check for duplicates then append ----
    existing_links = set()
    for row in ws_all.iter_rows(values_only=True):
        if row[13]:
            existing_links.add(str(row[13]))

    def safe_append(ws, row_tuple):
        link = row_tuple[13]
        if link and str(link) in existing_links:
            print(f'  SKIP (dupe): {row_tuple[2]}')
            return False
        ws.append(row_tuple)
        if link:
            existing_links.add(str(link))
        print(f'  ADD: {row_tuple[2]} → {ws.title}')
        return True

    print('=== Adding to 總覽 All Papers ===')
    safe_append(ws_all, GPAD_OVERVIEW)
    safe_append(ws_all, GAUGE_OVERVIEW)
    safe_append(ws_all, LOGFREE_OVERVIEW)
    safe_append(ws_all, LUMIN_OVERVIEW)

    print('=== Adding to MVTec 3D ===')
    # dupe-check per sheet
    existing_3d = set(str(r[13]) for r in ws_3d.iter_rows(values_only=True) if r[13])
    if GPAD_3D[13] not in existing_3d:
        ws_3d.append(GPAD_3D)
        print(f'  ADD: {GPAD_3D[2]} → MVTec 3D')
    else:
        print(f'  SKIP (dupe): {GPAD_3D[2]}')

    print('=== Adding to MVTec AD ===')
    existing_mvtec = set(str(r[13]) for r in ws_mvtec.iter_rows(values_only=True) if r[13])
    for row in [GAUGE_MVTEC, LUMIN_MVTEC]:
        if row[13] not in existing_mvtec:
            ws_mvtec.append(row)
            existing_mvtec.add(row[13])
            print(f'  ADD: {row[2]} → MVTec AD')
        else:
            print(f'  SKIP (dupe): {row[2]}')

    print('=== Adding to VisA ===')
    existing_visa = set(str(r[13]) for r in ws_visa.iter_rows(values_only=True) if r[13])
    for row in [GAUGE_VISA]:
        if row[13] not in existing_visa:
            ws_visa.append(row)
            existing_visa.add(row[13])
            print(f'  ADD: {row[2]} → VisA')
        else:
            print(f'  SKIP (dupe): {row[2]}')

    print('=== Adding to MVTec LOCO ===')
    existing_loco = set(str(r[13]) for r in ws_loco.iter_rows(values_only=True) if r[13])
    for row in [LOGFREE_LOCO]:
        if row[13] not in existing_loco:
            ws_loco.append(row)
            existing_loco.add(row[13])
            print(f'  ADD: {row[2]} → MVTec LOCO')
        else:
            print(f'  SKIP (dupe): {row[2]}')

    wb.save(str(AD_XLSX))
    print(f'\n✅ Saved: {AD_XLSX}')


if __name__ == '__main__':
    main()
