#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Journal harvest update — 2026-09-20
Adds verified new papers found in the weekly journal-targeted harvest.

Changes:
1. Add DriftAD (ACM MM 2026) — Few-shot, verified from GitHub README
2. Add TopoTTA (IEEE TPAMI 2026) — TTA/segmentation, ⏳ Pending
3. Update SubspaceAD (CVPR 2026) — Fill in verified metrics from GitHub README
4. Add PSMP-CLIP (arXiv 2609.16785) — Zero-shot CLIP, ⏳ Pending
"""

from pathlib import Path
import openpyxl

AD_XLSX = Path(__file__).resolve().parent.parent / 'AnomalyDetection_Papers_Summary_v10_20260425.xlsx'

# ---------------------------------------------------------------------------
# New row data
# ---------------------------------------------------------------------------
# Each entry: (dataset_col, category, method, authors, venue, date, status,
#              i_auroc, p_auroc, p_ap, p_pro, fps, notes, link, github)

DRIFTAD_OVERVIEW = (
    'MVTec AD, VisA (few-shot 1/2/4-shot, NON-STANDARD PROTOCOL)',
    '基於 VLM/CLIP (CLIP-based) / 少樣本 (Few-shot)',
    'DriftAD (Visually-Guided Text Drift for Few-Shot IAD)',
    'Wenyang Liu, Tianyi Liu, Dongshuo Zhang, Kejun Wu, Adams Wai-Kin Kong',
    'ACM MM 2026 / arXiv 2608.23723',
    '2026-08',
    '已發表',
    97.2,   # MVTec 1-shot I-AUROC (4-shot: 98.0)
    96.8,   # MVTec 1-shot P-AUROC (4-shot: 97.2)
    'N/A',
    92.2,   # MVTec 1-shot PRO (4-shot: 92.7)
    'N/A',
    '🟢 Few-shot verified — GitHub README (wenyang001/DriftAD); MVTec AD 1-shot: I-AUROC=97.2, P-AUROC=96.8, PRO=92.2; 4-shot: I-AUROC=98.0, P-AUROC=97.2, PRO=92.7; VisA 1-shot: I-AUROC=93.1, P-AUROC=97.4, PRO=86.7; 4-shot: I-AUROC=94.0, P-AUROC=98.0, PRO=88.3; OpenCLIP ViT-H/14; NON-STANDARD (few-shot 1/2/4-shot) — NOT comparable to standard unsupervised AD. ACM MM 2026 harvest 2026-09-20. [基於 VLM/CLIP / 少樣本]',
    'https://arxiv.org/abs/2608.23723',
    'https://github.com/wenyang001/DriftAD',
)

DRIFTAD_MVTEC = (
    'MVTec AD (few-shot 1-shot; NON-STANDARD PROTOCOL)',
    '基於 VLM/CLIP (CLIP-based) / 少樣本 (Few-shot)',
    'DriftAD (Visually-Guided Text Drift for Few-Shot IAD)',
    'Wenyang Liu, Tianyi Liu, Dongshuo Zhang, Kejun Wu, Adams Wai-Kin Kong',
    'ACM MM 2026 / arXiv 2608.23723',
    '2026-08',
    '已發表',
    97.2,
    96.8,
    'N/A',
    92.2,
    'N/A',
    '🟢 Few-shot verified — GitHub README (wenyang001/DriftAD); 1-shot I-AUROC=97.2, P-AUROC=96.8, PRO=92.2; 4-shot I-AUROC=98.0/97.2/92.7; NON-STANDARD: few-shot protocol, NOT comparable to standard unsupervised AD. [基於 VLM/CLIP / 少樣本]',
    'https://arxiv.org/abs/2608.23723',
    'https://github.com/wenyang001/DriftAD',
)

DRIFTAD_VISA = (
    'VisA (few-shot 1-shot; NON-STANDARD PROTOCOL)',
    '基於 VLM/CLIP (CLIP-based) / 少樣本 (Few-shot)',
    'DriftAD (Visually-Guided Text Drift for Few-Shot IAD)',
    'Wenyang Liu, Tianyi Liu, Dongshuo Zhang, Kejun Wu, Adams Wai-Kin Kong',
    'ACM MM 2026 / arXiv 2608.23723',
    '2026-08',
    '已發表',
    93.1,
    97.4,
    'N/A',
    86.7,
    'N/A',
    '🟢 Few-shot verified — GitHub README (wenyang001/DriftAD); 1-shot I-AUROC=93.1, P-AUROC=97.4, PRO=86.7; 4-shot I-AUROC=94.0/98.0/88.3; NON-STANDARD: few-shot protocol. [基於 VLM/CLIP / 少樣本]',
    'https://arxiv.org/abs/2608.23723',
    'https://github.com/wenyang001/DriftAD',
)

TOPOTTA_OVERVIEW = (
    'MVTec AD, VisA, Real-IAD, MVTec 3D-AD, MVTec LOCO',
    '基於表徵 (Representation) / 測試時適應 (Test-Time Adaptation) / 拓樸感知 (Topology-Aware)',
    'TopoTTA (Learning Topology-Aware Representations via TTA for Anomaly Segmentation)',
    'Ali Zia, Usman Ali, Abdul Rehman, Umer Ramzan, Kang Han, Muhammad Faheem, Shahnawaz Qureshi, Wei Xiang',
    'IEEE TPAMI 2026 (Early Access) / arXiv 2606.28268 / DOI 10.1109/TPAMI.2026.3710196 / PubMed 42397999',
    '2026-07',
    '⏳ Pending',
    None,
    None,
    None,
    None,
    None,
    '⏳ Pending — IEEE TPAMI 2026 Early Access (DOI 10.1109/TPAMI.2026.3710196, published 2026-07-03); arXiv 2606.28268; PubMed 42397999; arXiv HTML blocked — cannot directly verify Table I; search snippet cites backbone-dependent results: w/ Dinomaly: MVTec I-AUROC≈99.6/P-AUROC≈98.4/PRO≈94.8; VisA I-AUROC≈98.7/P-AUROC≈98.7/PRO≈94.5 (NOT standalone TopoTTA metrics — backbone-dependent TTA post-processing). Key claim: +15% avg F1 over SOTA across 6 benchmarks via persistent homology TTA. Project page: topotta.github.io. IEEE TPAMI journal harvest 2026-09-20. [基於表徵 / TTA / 拓樸感知]',
    'https://arxiv.org/abs/2606.28268',
    'https://topotta.github.io/',
)

TOPOTTA_MVTEC = (
    'MVTec AD (TTA post-processing; backbone-dependent; NON-STANDARD PROTOCOL)',
    '基於表徵 (Representation) / 測試時適應 (Test-Time Adaptation)',
    'TopoTTA (Learning Topology-Aware Representations via TTA for Anomaly Segmentation)',
    'Ali Zia, Usman Ali, Abdul Rehman, Umer Ramzan, Kang Han, Muhammad Faheem, Shahnawaz Qureshi, Wei Xiang',
    'IEEE TPAMI 2026 (Early Access) / arXiv 2606.28268',
    '2026-07',
    '⏳ Pending',
    None,
    None,
    None,
    None,
    None,
    '⏳ Pending — arXiv HTML blocked; search snippet cites w/ Dinomaly backbone: I-AUROC≈99.6/P-AUROC≈98.4/PRO≈94.8; backbone-dependent TTA — NOT standalone method. [TTA / 拓樸感知]',
    'https://arxiv.org/abs/2606.28268',
    'https://topotta.github.io/',
)

TOPOTTA_VISA = (
    'VisA (TTA post-processing; backbone-dependent; NON-STANDARD PROTOCOL)',
    '基於表徵 (Representation) / 測試時適應 (Test-Time Adaptation)',
    'TopoTTA (Learning Topology-Aware Representations via TTA for Anomaly Segmentation)',
    'Ali Zia, Usman Ali, Abdul Rehman, Umer Ramzan, Kang Han, Muhammad Faheem, Shahnawaz Qureshi, Wei Xiang',
    'IEEE TPAMI 2026 (Early Access) / arXiv 2606.28268',
    '2026-07',
    '⏳ Pending',
    None,
    None,
    None,
    None,
    None,
    '⏳ Pending — arXiv HTML blocked; search snippet cites w/ Dinomaly: I-AUROC≈98.7/P-AUROC≈98.7/PRO≈94.5; backbone-dependent TTA. [TTA / 拓樸感知]',
    'https://arxiv.org/abs/2606.28268',
    'https://topotta.github.io/',
)

PSMP_CLIP_OVERVIEW = (
    'MVTec AD, VisA, BTAD (cross-dataset zero-shot protocol)',
    '基於 VLM/CLIP (CLIP-based) / SAM輔助 (SAM-assisted) / 零樣本 (Zero-shot)',
    'PSMP-CLIP (Patch-Prompt SAM and Multi-Semantic Prompting for CLIP-Based Zero-Shot AD)',
    'see arXiv 2609.16785',
    'arXiv 2609.16785 (2026-09-15)',
    '2026-09',
    '⏳ Pending',
    None,
    None,
    None,
    None,
    None,
    '⏳ Pending — arXiv 2609.16785 (Sept 15, 2026); zero-shot cross-dataset protocol (trained on VisA→evaluated on MVTec AD and others); search snippet reports MVTec AD P-AUROC≈93.7 (cross-dataset zero-shot — NOT comparable to standard unsupervised AD); ViT-L-14-336 CLIP + SAM2.1; 14 datasets; no GitHub found as of 2026-09-20. NON-STANDARD (cross-dataset zero-shot). [基於 VLM/CLIP / SAM輔助 / 零樣本]',
    'https://arxiv.org/abs/2609.16785',
    None,
)


def get_overview_row_index(ws, method_name):
    """Find the row index of an existing method in the overview sheet."""
    for i, row in enumerate(ws.iter_rows(min_row=4), start=4):
        if row[2].value and method_name.lower() in str(row[2].value).lower():
            return i
    return None


def append_row(ws, row_data):
    """Append a row of data to the given sheet."""
    ws.append(list(row_data))


def update_row(ws, row_idx, row_data):
    """Update specific columns of an existing row."""
    row_data_list = list(row_data)
    for col_idx, val in enumerate(row_data_list, start=1):
        ws.cell(row=row_idx, column=col_idx).value = val


def main():
    wb = openpyxl.load_workbook(AD_XLSX)

    overview_ws = wb['總覽 All Papers']
    mvtec_ws = wb['MVTec AD']
    visa_ws = wb['VisA']

    # ------------------------------------------------------------------
    # 1. Update SubspaceAD in overview with verified metrics
    # ------------------------------------------------------------------
    subspace_row = get_overview_row_index(overview_ws, 'SubspaceAD')
    if subspace_row:
        print(f'Updating SubspaceAD at row {subspace_row}')
        ws_row = overview_ws[subspace_row]
        # Col A (dataset)
        overview_ws.cell(row=subspace_row, column=1).value = 'MVTec AD, VisA (few-shot 1-shot; NON-STANDARD PROTOCOL)'
        # Col H (I-AUROC)
        overview_ws.cell(row=subspace_row, column=8).value = 97.1
        # Col I (P-AUROC)
        overview_ws.cell(row=subspace_row, column=9).value = 97.5
        # Col J (P-AP) — N/A
        overview_ws.cell(row=subspace_row, column=10).value = 'N/A'
        # Col K (P-PRO) — N/A (not reported)
        overview_ws.cell(row=subspace_row, column=11).value = 'N/A'
        # Col M (notes)
        overview_ws.cell(row=subspace_row, column=13).value = (
            '🟢 Few-shot verified — GitHub README (CLendering/SubspaceAD); '
            'MVTec AD 1-shot: I-AUROC=97.1, P-AUROC=97.5; '
            'VisA 1-shot: I-AUROC=93.2, P-AUROC=98.2; '
            'DINOv2-based training-free PCA subspace method; CVPR 2026. '
            'NON-STANDARD (few-shot 1-shot) — NOT comparable to standard unsupervised AD. '
            'Journal harvest update 2026-09-20. [基於表徵 / 少樣本 / Training-free]'
        )
        print(f'  → SubspaceAD updated: I-AUROC=97.1, P-AUROC=97.5 (MVTec AD 1-shot)')
    else:
        print('SubspaceAD not found in overview — skipping update')

    # ------------------------------------------------------------------
    # 2. Add DriftAD
    # ------------------------------------------------------------------
    # Dupe check
    if get_overview_row_index(overview_ws, 'DriftAD'):
        print('DriftAD already in overview — skipping')
    else:
        print('Adding DriftAD to 總覽...')
        append_row(overview_ws, DRIFTAD_OVERVIEW)

    # Add to MVTec AD sheet
    found_mvtec = False
    for row in mvtec_ws.iter_rows(min_row=4, values_only=True):
        if row[2] and 'DriftAD' in str(row[2]):
            found_mvtec = True
            break
    if not found_mvtec:
        print('Adding DriftAD to MVTec AD sheet...')
        append_row(mvtec_ws, DRIFTAD_MVTEC)

    # Add to VisA sheet
    found_visa = False
    for row in visa_ws.iter_rows(min_row=4, values_only=True):
        if row[2] and 'DriftAD' in str(row[2]):
            found_visa = True
            break
    if not found_visa:
        print('Adding DriftAD to VisA sheet...')
        append_row(visa_ws, DRIFTAD_VISA)

    # ------------------------------------------------------------------
    # 3. Add TopoTTA
    # ------------------------------------------------------------------
    if get_overview_row_index(overview_ws, 'TopoTTA'):
        print('TopoTTA already in overview — skipping')
    else:
        print('Adding TopoTTA to 總覽...')
        append_row(overview_ws, TOPOTTA_OVERVIEW)

    # Add to MVTec AD sheet
    found_mvtec_topo = False
    for row in mvtec_ws.iter_rows(min_row=4, values_only=True):
        if row[2] and 'TopoTTA' in str(row[2]):
            found_mvtec_topo = True
            break
    if not found_mvtec_topo:
        print('Adding TopoTTA to MVTec AD sheet...')
        append_row(mvtec_ws, TOPOTTA_MVTEC)

    # Add to VisA sheet
    found_visa_topo = False
    for row in visa_ws.iter_rows(min_row=4, values_only=True):
        if row[2] and 'TopoTTA' in str(row[2]):
            found_visa_topo = True
            break
    if not found_visa_topo:
        print('Adding TopoTTA to VisA sheet...')
        append_row(visa_ws, TOPOTTA_VISA)

    # ------------------------------------------------------------------
    # 4. Add PSMP-CLIP
    # ------------------------------------------------------------------
    if get_overview_row_index(overview_ws, 'PSMP-CLIP'):
        print('PSMP-CLIP already in overview — skipping')
    else:
        print('Adding PSMP-CLIP to 總覽...')
        append_row(overview_ws, PSMP_CLIP_OVERVIEW)

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------
    wb.save(AD_XLSX)
    print(f'\nSaved {AD_XLSX}')
    print('\nSummary of changes:')
    print('  - SubspaceAD: metrics updated (1-shot, few-shot verified)')
    print('  - DriftAD: added to 總覽, MVTec AD, VisA (ACM MM 2026, few-shot verified)')
    print('  - TopoTTA: added to 總覽, MVTec AD, VisA (IEEE TPAMI 2026, pending)')
    print('  - PSMP-CLIP: added to 總覽 (arXiv preprint, zero-shot, pending)')


if __name__ == '__main__':
    main()
