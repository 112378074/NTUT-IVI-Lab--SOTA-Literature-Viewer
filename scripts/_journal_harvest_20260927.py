#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Journal harvest update — 2026-09-27
Weekly journal-targeted harvest; adds new AD candidates found this week.

Search scope: IEEE TPAMI, IEEE TII, IEEE TNNLS, Pattern Recognition, IJCV,
  IEEE TIP, IEEE TCSVT, Computers in Industry, Information Fusion,
  Knowledge-Based Systems, Information Sciences, Neural Networks.

Findings this run:
  — No new CONFIRMED journal papers found in target journals (past 7 days).
  — arXiv HTML blocked; no GitHub repos found for new candidates.
  — 5 arXiv preprints staged as ⏳ Pending:
    1. SwinAD (2607.14534, Jul 2026) — standard unsupervised AD, MVTec AD/VisA/Real-IAD
    2. MambaADv2 (2606.23126, Jun 2026) — standard unsupervised AD, 6 benchmarks incl. MVTec 3D
    3. NC-TFAD (2609.03406, Sep 2026) — continual AD (NON-STANDARD), MVTec AD/VisA
    4. PL-SCEA (2609.03655, Sep 2026) — few-shot AD (NON-STANDARD), MVTec AD/VisA
    5. A Principled Approach (2609.21800, Sep 2026) — standard unsupervised AD, MVTec AD
"""

from pathlib import Path
import openpyxl

AD_XLSX = Path(__file__).resolve().parent.parent / 'AnomalyDetection_Papers_Summary_v10_20260425.xlsx'

# ---------------------------------------------------------------------------
# Row data — column order matches xlsx headers:
# (資料集, 類別, 方法, 作者, 發表, 年月, 狀態,
#  I-AUROC, P-AUROC, P-AP, P-PRO, FPS, 備註, 連結, GitHub)
# ---------------------------------------------------------------------------

SWINAD_OVERVIEW = (
    'MVTec AD, VisA, Real-IAD',
    '基於重構 (Reconstruction-based) / Swin Transformer V2',
    'SwinAD (Multi-stage Feature Reconstruction for Unsupervised Industrial AD)',
    'Huong Ninh, Chien Thai, Mai Xuan Trang, Vu-Minh Le, Thanh Ha Le, Long Tran',
    'arXiv 2607.14534 (2026-07-16)',
    '2026-07',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2607.14534 (Jul 16, 2026); arXiv HTML blocked, no GitHub repo found; '
    'web-search snippet from arXiv HTML: MVTec AD I-AUROC≈99.4%, I-AP≈99.7% (not directly '
    'verified from authoritative table). Frozen Swin Transformer V2 encoder + '
    'feature-diversity-preserving reconstruction decoder; stage-wise bottleneck dropout; '
    'MVTec AD, VisA, Real-IAD benchmarks. No journal affiliation confirmed. '
    'Journal harvest 2026-09-27. [基於重構 / Swin Transformer]',
    'https://arxiv.org/abs/2607.14534',
    None,
)

SWINAD_MVTEC = (
    'MVTec AD',
    '基於重構 (Reconstruction-based) / Swin Transformer V2',
    'SwinAD (Multi-stage Feature Reconstruction for Unsupervised Industrial AD)',
    'Huong Ninh, Chien Thai, Mai Xuan Trang, Vu-Minh Le, Thanh Ha Le, Long Tran',
    'arXiv 2607.14534 (2026-07-16)',
    '2026-07',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-search snippet: I-AUROC≈99.4%, I-AP≈99.7% '
    '(not directly verified). [基於重構]',
    'https://arxiv.org/abs/2607.14534',
    None,
)

SWINAD_VISA = (
    'VisA',
    '基於重構 (Reconstruction-based) / Swin Transformer V2',
    'SwinAD (Multi-stage Feature Reconstruction for Unsupervised Industrial AD)',
    'Huong Ninh, Chien Thai, Mai Xuan Trang, Vu-Minh Le, Thanh Ha Le, Long Tran',
    'arXiv 2607.14534 (2026-07-16)',
    '2026-07',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; no specific VisA metrics in search snippets. [基於重構]',
    'https://arxiv.org/abs/2607.14534',
    None,
)

MAMBAADV2_OVERVIEW = (
    'MVTec AD, VisA, Real-IAD, MVTec 3D-AD, COCO-AD, Uni-Medical',
    '基於 Mamba/SSM (State Space Model-based) / 重構 (Reconstruction)',
    'MambaADv2 (Evolving Duality-enhanced State Space Model for Unsupervised AD)',
    'see arXiv 2606.23126',
    'arXiv 2606.23126 (2026-06-22)',
    '2026-06',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2606.23126 (Jun 22, 2026); arXiv HTML blocked, no GitHub repo found; '
    'web-search snippet: MVTec AD I-AUROC≈99.2%, mAD=87.2 (composite metric), VisA mAD=80.1; '
    'not directly verified from authoritative table. Duality-enhanced State Space (DSS) modules; '
    '6 benchmarks: MVTec AD, VisA, Real-IAD, MVTec 3D-AD, COCO-AD, Uni-Medical. '
    'No journal affiliation confirmed. Journal harvest 2026-09-27. [基於 Mamba/SSM / 重構]',
    'https://arxiv.org/abs/2606.23126',
    None,
)

MAMBAADV2_MVTEC = (
    'MVTec AD',
    '基於 Mamba/SSM (State Space Model-based) / 重構 (Reconstruction)',
    'MambaADv2 (Evolving Duality-enhanced State Space Model for Unsupervised AD)',
    'see arXiv 2606.23126',
    'arXiv 2606.23126 (2026-06-22)',
    '2026-06',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-search snippet: I-AUROC≈99.2%, mAD=87.2 '
    '(mAD is composite, not standard I-AUROC). [基於 Mamba/SSM]',
    'https://arxiv.org/abs/2606.23126',
    None,
)

MAMBAADV2_VISA = (
    'VisA',
    '基於 Mamba/SSM (State Space Model-based) / 重構 (Reconstruction)',
    'MambaADv2 (Evolving Duality-enhanced State Space Model for Unsupervised AD)',
    'see arXiv 2606.23126',
    'arXiv 2606.23126 (2026-06-22)',
    '2026-06',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-search snippet: VisA mAD=80.1 '
    '(composite metric, not standard I-AUROC). [基於 Mamba/SSM]',
    'https://arxiv.org/abs/2606.23126',
    None,
)

MAMBAADV2_3D = (
    'MVTec 3D-AD',
    '基於 Mamba/SSM (State Space Model-based) / 重構 (Reconstruction)',
    'MambaADv2 (Evolving Duality-enhanced State Space Model for Unsupervised AD)',
    'see arXiv 2606.23126',
    'arXiv 2606.23126 (2026-06-22)',
    '2026-06',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; no specific MVTec 3D metrics in search snippets. [基於 Mamba/SSM]',
    'https://arxiv.org/abs/2606.23126',
    None,
)

NCTFAD_OVERVIEW = (
    'MVTec AD, VisA (task-free continual learning protocol; NON-STANDARD PROTOCOL)',
    '持續學習 (Continual Learning) / 神經坍縮引導 (Neural-Collapse-guided) — NON-STANDARD PROTOCOL',
    'NC-TFAD (Neural-Collapse-guided Task-Free Continual Anomaly Detection)',
    'Xiaotong Kong, Chaoyang Song, Ziai Zhou, Jinxia Zhang, Kanjian Zhang, Haikun Wei (Southeast University)',
    'arXiv 2609.03406 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.03406 (Sep 3, 2026); arXiv HTML blocked, no GitHub repo found; '
    'web-search snippet: MVTec AD I-AUROC≈82.1% in task-free continual learning setting '
    '(NOT comparable to standard unsupervised AD — continual protocol). '
    'Neural-collapse-inspired geometry-driven framework; Focal Neural Collapse Contrastive (FNCC) loss; '
    'MVTec AD and VisA datasets under task-free continual learning protocol. '
    'NON-STANDARD: continual learning protocol — NOT comparable to standard unsupervised AD. '
    'No journal affiliation confirmed. Journal harvest 2026-09-27. [持續學習 / 神經坍縮引導]',
    'https://arxiv.org/abs/2609.03406',
    None,
)

NCTFAD_MVTEC = (
    'MVTec AD (task-free continual learning protocol; NON-STANDARD PROTOCOL)',
    '持續學習 (Continual Learning) / 神經坍縮引導 (Neural-Collapse-guided) — NON-STANDARD PROTOCOL',
    'NC-TFAD (Neural-Collapse-guided Task-Free Continual Anomaly Detection)',
    'Xiaotong Kong, Chaoyang Song, Ziai Zhou, Jinxia Zhang, Kanjian Zhang, Haikun Wei (Southeast University)',
    'arXiv 2609.03406 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; web-search snippet: I-AUROC≈82.1% (continual setting). '
    'NON-STANDARD: task-free continual learning — NOT comparable to standard unsupervised AD. [持續學習]',
    'https://arxiv.org/abs/2609.03406',
    None,
)

NCTFAD_VISA = (
    'VisA (task-free continual learning protocol; NON-STANDARD PROTOCOL)',
    '持續學習 (Continual Learning) / 神經坍縮引導 (Neural-Collapse-guided) — NON-STANDARD PROTOCOL',
    'NC-TFAD (Neural-Collapse-guided Task-Free Continual Anomaly Detection)',
    'Xiaotong Kong, Chaoyang Song, Ziai Zhou, Jinxia Zhang, Kanjian Zhang, Haikun Wei (Southeast University)',
    'arXiv 2609.03406 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; no specific VisA metrics in snippets (continual protocol). '
    'NON-STANDARD: task-free continual learning. [持續學習]',
    'https://arxiv.org/abs/2609.03406',
    None,
)

PLSCEA_OVERVIEW = (
    'MVTec AD, VisA (few-shot; NON-STANDARD PROTOCOL)',
    '基於 VFM/基礎模型 (Vision Foundation Model) / 少樣本 (Few-shot) — NON-STANDARD PROTOCOL',
    'PL-SCEA (Power-Law Self-Correlation Enhanced Attention for Few-Shot Industrial AD)',
    'Xiaoyu Yang, Qixing Wu, Huixian Zhao, Changlong Jin',
    'arXiv 2609.03655 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.03655 (Sep 3, 2026); arXiv HTML blocked, no GitHub repo found; '
    'no specific metrics in search snippets. Reconfigures pretrained attention of VFMs for few-shot '
    'industrial AD; Power-Law Self-Correlation Enhanced Attention (PL-SCEA); token-adaptive '
    'self-correlations over contextualized value features; MVTec AD and VisA under few-shot settings. '
    'NON-STANDARD: few-shot protocol — NOT comparable to standard unsupervised AD. '
    'No journal affiliation confirmed. Journal harvest 2026-09-27. [VFM / 少樣本]',
    'https://arxiv.org/abs/2609.03655',
    None,
)

PLSCEA_MVTEC = (
    'MVTec AD (few-shot; NON-STANDARD PROTOCOL)',
    '基於 VFM/基礎模型 (Vision Foundation Model) / 少樣本 (Few-shot) — NON-STANDARD PROTOCOL',
    'PL-SCEA (Power-Law Self-Correlation Enhanced Attention for Few-Shot Industrial AD)',
    'Xiaoyu Yang, Qixing Wu, Huixian Zhao, Changlong Jin',
    'arXiv 2609.03655 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; no specific metrics found. '
    'NON-STANDARD: few-shot protocol. [VFM / 少樣本]',
    'https://arxiv.org/abs/2609.03655',
    None,
)

PLSCEA_VISA = (
    'VisA (few-shot; NON-STANDARD PROTOCOL)',
    '基於 VFM/基礎模型 (Vision Foundation Model) / 少樣本 (Few-shot) — NON-STANDARD PROTOCOL',
    'PL-SCEA (Power-Law Self-Correlation Enhanced Attention for Few-Shot Industrial AD)',
    'Xiaoyu Yang, Qixing Wu, Huixian Zhao, Changlong Jin',
    'arXiv 2609.03655 (2026-09-03)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; no specific VisA metrics. NON-STANDARD: few-shot. [VFM / 少樣本]',
    'https://arxiv.org/abs/2609.03655',
    None,
)

PRINCIPLED_OVERVIEW = (
    'MVTec AD (object-class subset)',
    '基於理論/貝葉斯 (Bayesian/Theoretical) / 無監督 AD',
    'PrincipalUAD (A Principled Approach to Unsupervised Anomaly Detection)',
    'James Myles, Matthew Baugh, Johanna P. Müller, Bernhard Kainz, Yingzhen Li',
    'arXiv 2609.21800 (2026-09-18)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv 2609.21800 (Sep 18, 2026); arXiv HTML blocked, no GitHub repo found; '
    'web-search snippet: improves MVTec AD object-class AUROC by 2.3% via adapting the '
    'underlying corruption model (exact absolute AUROC not confirmed from authoritative table). '
    'Reformulates UAD as Bayesian inverse problem; probabilistic anomaly score as energy of '
    'inferred corruption parameters; derives existing methods as special cases. '
    'Evaluated on MVTec AD (object-class subset). No journal affiliation confirmed. '
    'Journal harvest 2026-09-27. [基於理論/貝葉斯]',
    'https://arxiv.org/abs/2609.21800',
    None,
)

PRINCIPLED_MVTEC = (
    'MVTec AD (object-class subset)',
    '基於理論/貝葉斯 (Bayesian/Theoretical) / 無監督 AD',
    'PrincipalUAD (A Principled Approach to Unsupervised Anomaly Detection)',
    'James Myles, Matthew Baugh, Johanna P. Müller, Bernhard Kainz, Yingzhen Li',
    'arXiv 2609.21800 (2026-09-18)',
    '2026-09',
    '⏳ Pending',
    None, None, None, None, None,
    '⏳ Pending — arXiv HTML blocked; snippet: +2.3% object-class AUROC (absolute value not confirmed). '
    'Bayesian inverse problem formulation. [基於理論/貝葉斯]',
    'https://arxiv.org/abs/2609.21800',
    None,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_method_row(ws, method_key):
    for i, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if row[2] and method_key.lower() in str(row[2]).lower():
            return i
    return None


def dupe_check(ws, method_key, min_row=2):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if row[2] and method_key.lower() in str(row[2]).lower():
            return True
    return False


def append_row(ws, row_data):
    ws.append(list(row_data))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    wb = openpyxl.load_workbook(AD_XLSX)
    overview_ws = wb['總覽 All Papers']
    mvtec_ws = wb['MVTec AD']
    visa_ws = wb['VisA']
    mvtec3d_ws = wb['MVTec 3D']

    added = []

    # ------------------------------------------------------------------
    # 1. SwinAD
    # ------------------------------------------------------------------
    if not dupe_check(overview_ws, 'SwinAD'):
        print('Adding SwinAD to 總覽...')
        append_row(overview_ws, SWINAD_OVERVIEW)
        added.append('SwinAD → 總覽')
    else:
        print('SwinAD already in 總覽 — skipping')

    if not dupe_check(mvtec_ws, 'SwinAD', min_row=4):
        print('Adding SwinAD to MVTec AD sheet...')
        append_row(mvtec_ws, SWINAD_MVTEC)
        added.append('SwinAD → MVTec AD')
    else:
        print('SwinAD already in MVTec AD — skipping')

    if not dupe_check(visa_ws, 'SwinAD', min_row=4):
        print('Adding SwinAD to VisA sheet...')
        append_row(visa_ws, SWINAD_VISA)
        added.append('SwinAD → VisA')
    else:
        print('SwinAD already in VisA — skipping')

    # ------------------------------------------------------------------
    # 2. MambaADv2
    # ------------------------------------------------------------------
    if not dupe_check(overview_ws, 'MambaADv2'):
        print('Adding MambaADv2 to 總覽...')
        append_row(overview_ws, MAMBAADV2_OVERVIEW)
        added.append('MambaADv2 → 總覽')
    else:
        print('MambaADv2 already in 總覽 — skipping')

    if not dupe_check(mvtec_ws, 'MambaADv2', min_row=4):
        print('Adding MambaADv2 to MVTec AD sheet...')
        append_row(mvtec_ws, MAMBAADV2_MVTEC)
        added.append('MambaADv2 → MVTec AD')
    else:
        print('MambaADv2 already in MVTec AD — skipping')

    if not dupe_check(visa_ws, 'MambaADv2', min_row=4):
        print('Adding MambaADv2 to VisA sheet...')
        append_row(visa_ws, MAMBAADV2_VISA)
        added.append('MambaADv2 → VisA')
    else:
        print('MambaADv2 already in VisA — skipping')

    if not dupe_check(mvtec3d_ws, 'MambaADv2', min_row=4):
        print('Adding MambaADv2 to MVTec 3D sheet...')
        append_row(mvtec3d_ws, MAMBAADV2_3D)
        added.append('MambaADv2 → MVTec 3D')
    else:
        print('MambaADv2 already in MVTec 3D — skipping')

    # ------------------------------------------------------------------
    # 3. NC-TFAD
    # ------------------------------------------------------------------
    if not dupe_check(overview_ws, 'NC-TFAD'):
        print('Adding NC-TFAD to 總覽...')
        append_row(overview_ws, NCTFAD_OVERVIEW)
        added.append('NC-TFAD → 總覽')
    else:
        print('NC-TFAD already in 總覽 — skipping')

    if not dupe_check(mvtec_ws, 'NC-TFAD', min_row=4):
        print('Adding NC-TFAD to MVTec AD sheet...')
        append_row(mvtec_ws, NCTFAD_MVTEC)
        added.append('NC-TFAD → MVTec AD (continual, non-standard)')
    else:
        print('NC-TFAD already in MVTec AD — skipping')

    if not dupe_check(visa_ws, 'NC-TFAD', min_row=4):
        print('Adding NC-TFAD to VisA sheet...')
        append_row(visa_ws, NCTFAD_VISA)
        added.append('NC-TFAD → VisA (continual, non-standard)')
    else:
        print('NC-TFAD already in VisA — skipping')

    # ------------------------------------------------------------------
    # 4. PL-SCEA
    # ------------------------------------------------------------------
    if not dupe_check(overview_ws, 'PL-SCEA'):
        print('Adding PL-SCEA to 總覽...')
        append_row(overview_ws, PLSCEA_OVERVIEW)
        added.append('PL-SCEA → 總覽')
    else:
        print('PL-SCEA already in 總覽 — skipping')

    if not dupe_check(mvtec_ws, 'PL-SCEA', min_row=4):
        print('Adding PL-SCEA to MVTec AD sheet...')
        append_row(mvtec_ws, PLSCEA_MVTEC)
        added.append('PL-SCEA → MVTec AD (few-shot, non-standard)')
    else:
        print('PL-SCEA already in MVTec AD — skipping')

    if not dupe_check(visa_ws, 'PL-SCEA', min_row=4):
        print('Adding PL-SCEA to VisA sheet...')
        append_row(visa_ws, PLSCEA_VISA)
        added.append('PL-SCEA → VisA (few-shot, non-standard)')
    else:
        print('PL-SCEA already in VisA — skipping')

    # ------------------------------------------------------------------
    # 5. A Principled Approach to UAD (PrincipalUAD)
    # ------------------------------------------------------------------
    if not dupe_check(overview_ws, 'PrincipalUAD'):
        print('Adding PrincipalUAD to 總覽...')
        append_row(overview_ws, PRINCIPLED_OVERVIEW)
        added.append('PrincipalUAD → 總覽')
    else:
        print('PrincipalUAD already in 總覽 — skipping')

    if not dupe_check(mvtec_ws, 'PrincipalUAD', min_row=4):
        print('Adding PrincipalUAD to MVTec AD sheet...')
        append_row(mvtec_ws, PRINCIPLED_MVTEC)
        added.append('PrincipalUAD → MVTec AD')
    else:
        print('PrincipalUAD already in MVTec AD — skipping')

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------
    wb.save(AD_XLSX)
    print(f'\nSaved {AD_XLSX}')
    print('\nSummary of changes:')
    for item in added:
        print(f'  + {item}')
    print(f'\nTotal rows added: {len(added)}')
    print('\nAll entries staged as ⏳ Pending — arXiv HTML blocked; no GitHub repos found;')
    print('re-verify when repos/HTML become accessible.')


if __name__ == '__main__':
    main()
